import re
import zipfile
from html.parser import HTMLParser

import pytest
from openpyxl import load_workbook

from inflationarg import formato, site

PAGINAS = [
    "index.html",
    "dashboard/index.html",
    "calculadora/index.html",
    "sql/index.html",
    "metodologia/index.html",
    "sobre-mi/index.html",
    "404.html",
]


# --------------------------------------------------------------------------- JSON con esquema

ESQUEMA = {
    "meta": {"fecha_descarga": str, "fecha_corte": str, "base_analisis": str, "series": int, "observaciones": int},
    "kpis": dict,
    "sectores": dict,
    "regiones": dict,
    "graficos": {
        "salario_real": dict,
        "inflacion": dict,
        "categorias": dict,
        "regiones": dict,
        "poder_compra": dict,
        "usd": dict,
    },
    "cobertura": list,
    "consultas": dict,
    "qa": {"fecha_descarga": str, "fecha_corte": str, "conteo": dict, "resultados": list, "marcas": list},
}


def _validar(valor, esquema, ruta="$"):
    if isinstance(esquema, dict):
        assert isinstance(valor, dict), ruta
        assert set(esquema) <= set(valor), f"{ruta}: faltan {set(esquema) - set(valor)}"
        for k, sub in esquema.items():
            _validar(valor[k], sub, f"{ruta}.{k}")
    else:
        assert isinstance(valor, esquema), f"{ruta}: se esperaba {esquema.__name__}"


def test_json_del_dashboard_cumple_el_esquema(dashboard):
    _validar(dashboard, ESQUEMA)
    for nombre, g in dashboard["graficos"].items():
        for serie in g["series"].values():
            assert len(serie) == len(g["fechas"]), nombre
    for nombre, c in dashboard["consultas"].items():
        assert set(c) == {"descripcion", "sql", "columnas", "filas"}, nombre
        for fila in c["filas"]:
            assert list(fila) == c["columnas"], nombre


def test_json_sin_nan(corrida):
    texto = (corrida.salida / "datos" / "dashboard.json").read_text(encoding="utf-8")
    assert "NaN" not in texto and "Infinity" not in texto


# --------------------------------------------------------------------------- Excel y Power BI


def test_excel_una_hoja_por_pregunta_y_grafico(corrida):
    wb = load_workbook(corrida.salida / "inflationarg.xlsx")
    assert wb.sheetnames[0] == "Leeme"
    assert len([h for h in wb.sheetnames if h[0].isdigit()]) == 9
    hoja = wb["Salario real mensual"]
    assert len(hoja._charts) == 1
    assert hoja.tables  # tabla con formato
    hoja = wb["1 Respuesta"]
    assert hoja["A5"].value == "Serie"  # encabezados legibles, sin columnas técnicas
    assert hoja["A2"].value.startswith("¿Le ganaron los salarios")
    assert hoja.page_setup.orientation == "landscape"
    leeme = wb["Leeme"]
    assert leeme["A5"].hyperlink.location == "'1 Respuesta'!A1"


def test_power_bi_modelo_estrella(corrida):
    archivos = sorted(p.name for p in (corrida.salida / "powerbi").glob("*.csv"))
    assert archivos == ["dim_fecha.csv", "dim_serie.csv", "fact_valor.csv", "salario_real.csv"]
    cabecera = (corrida.salida / "powerbi" / "fact_valor.csv").read_text(encoding="utf-8-sig").splitlines()[0]
    assert cabecera == "fecha,serie_id,valor,estado"


# --------------------------------------------------------------------------- Sitio


class Enlaces(HTMLParser):
    def __init__(self):
        super().__init__()
        self.refs, self.ids, self.imgs_sin_alt = [], set(), 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if "id" in a:
            self.ids.add(a["id"])
        for k in ("href", "src"):
            if a.get(k):
                self.refs.append(a[k])
        if tag == "img" and "alt" not in a:
            self.imgs_sin_alt += 1


@pytest.fixture(scope="module")
def sitio(corrida):
    return corrida.salida / "site"


def test_todas_las_paginas_renderizan(sitio):
    for pagina in PAGINAS:
        html = (sitio / pagina).read_text(encoding="utf-8")
        assert html.startswith("<!doctype html>"), pagina
        assert '<html lang="es-AR">' in html
        assert re.search(r"<title>[^<]+</title>", html)
        assert 'property="og:image"' in html
        visible = re.sub(r"<script.*?</script>", "", html, flags=re.S)
        for basura in (r"\bNone\b", r"\bnan\b", r"\{\{", r"\}\}", r"\bundefined\b"):
            assert not re.search(basura, visible), (pagina, basura)


def test_links_internos_validos(sitio):
    for pagina in PAGINAS:
        parser = Enlaces()
        parser.feed((sitio / pagina).read_text(encoding="utf-8"))
        assert parser.imgs_sin_alt == 0
        for ref in parser.refs:
            if ref.startswith("#"):
                assert ref[1:] in parser.ids, (pagina, ref)
                continue
            if not ref.startswith("/InflationArg/"):
                assert ref.startswith(("https://", "mailto:")), (pagina, ref)
                continue
            ruta, _, ancla = ref.removeprefix("/InflationArg/").partition("#")
            destino = sitio / ruta
            if ruta == "" or ruta.endswith("/"):
                destino = destino / "index.html"
            assert destino.exists(), (pagina, ref)
            if ancla:
                otro = Enlaces()
                otro.feed(destino.read_text(encoding="utf-8"))
                assert ancla in otro.ids, (pagina, ref)


def test_inicio_responde_la_pregunta_con_datos_de_la_base(sitio, dashboard):
    html = (sitio / "index.html").read_text(encoding="utf-8")
    total = dashboard["kpis"]["salario_real_total"]
    assert "<strong>No.</strong>" in html
    assert formato.porcentaje(abs(total["var_real_pct"]), 1) in html
    assert "julio de 2026" in html


def test_sitemap_robots_descargas(sitio):
    sitemap = (sitio / "sitemap.xml").read_text(encoding="utf-8")
    assert sitemap.count("<url>") == len(site.PAGINAS)
    assert "https://juanmbf.github.io/InflationArg/dashboard/" in sitemap
    assert "Sitemap:" in (sitio / "robots.txt").read_text(encoding="utf-8")
    with zipfile.ZipFile(sitio / "descargas" / "powerbi.zip") as zf:
        assert "fact_valor.csv" in zf.namelist()
    assert (sitio / "static" / "img" / "og.png").exists()


def test_echarts_con_version_fija_y_sri(sitio):
    html = (sitio / "dashboard" / "index.html").read_text(encoding="utf-8")
    assert "echarts@6.0.0/" in html
    assert 'integrity="sha384-' in html


def test_json_embebido_no_rompe_el_html():
    assert str(site._json_embebido({"a": "</script>"})) == '{"a":"\\u003c/script\\u003e"}'


def test_resaltado_sql():
    html = str(site.resaltar_sql("SELECT LAG(x) -- comentario <b>\nFROM t WHERE y = 'a'"))
    assert '<span class="k">SELECT</span>' in html
    assert '<span class="f">LAG</span>' in html
    assert '<span class="c">-- comentario &lt;b&gt;</span>' in html
    assert '<span class="s">&#x27;a&#x27;</span>' in html


def test_celdas_de_tabla():
    assert site.celda(-21.789, "var_real_pct") == "-21,79"
    assert site.celda(1442.0, "ripte_usd") == "1.442"
    assert site.celda("2026-07-01", "vale_hasta") == "jul-2026"
    assert site.celda(None, "x") == "–"
    assert site.celda("Salarios – sector público", "nombre") == "Sector público"
    assert site.clase_celda(-1.0, "var_real_pct") == "num neg"
