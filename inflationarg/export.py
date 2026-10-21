"""Exportaciones: JSON para el sitio, Excel con formato y CSV para Power BI.

Todo lo que se muestra sale de la base: vistas y consultas de sql/consultas.
"""

from __future__ import annotations

import json
import math
import re
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import config, db, formato
from .calculadora import Indices
from .formato import mes
from .qa import Reporte

NOMBRES_SECTOR = {
    "sal_privado": "Privado registrado",
    "sal_publico": "Público",
    "sal_no_registrado": "Privado no registrado",
    "sal_total": "Total",
}


def _limpio(valor):
    if isinstance(valor, float) and (math.isnan(valor) or math.isinf(valor)):
        return None
    return valor


def _registros(df: pd.DataFrame) -> list[dict]:
    return [{k: _limpio(v) for k, v in fila.items()} for fila in df.to_dict(orient="records")]


def _serie_ancha(df: pd.DataFrame, columna: str, valores: str, redondeo: int = 4) -> dict:
    """De formato largo a {"fechas": [...], "series": {id: [...]}} para ECharts."""
    ancho = df.pivot(index="fecha", columns=columna, values=valores).sort_index()
    return {
        "fechas": [f[:7] for f in ancho.index],
        "series": {c: [None if pd.isna(v) else round(float(v), redondeo) for v in ancho[c]] for c in ancho.columns},
    }


def resultados_consultas(con: sqlite3.Connection) -> dict[str, dict]:
    salida = {}
    for nombre, sql in db.consultas_de_negocio().items():
        df = db.consultar(con, sql)
        salida[nombre] = {
            "descripcion": db.descripcion_consulta(sql),
            "sql": sql,
            "columnas": list(df.columns),
            "filas": _registros(df),
        }
    return salida


def datos_dashboard(con: sqlite3.Connection, consultas: dict, reporte: Reporte) -> dict:
    q = lambda sql: db.consultar(con, sql)  # noqa: E731

    parametros = dict(con.execute("SELECT clave, valor FROM parametro").fetchall())
    real = q("SELECT fecha, serie_id, indice_real FROM v_salario_real")
    inflacion = q(
        "SELECT fecha, var_mensual * 100 AS mensual, var_interanual * 100 AS interanual "
        "FROM v_variacion WHERE serie_id = 'ipc_nacional' AND var_mensual IS NOT NULL ORDER BY fecha"
    )
    categorias = q(
        "SELECT fecha, serie_id, var_interanual * 100 AS v FROM v_variacion "
        "WHERE grupo IN ('ipc_categoria') OR serie_id = 'ipc_nacional'"
    )
    regiones = q("SELECT fecha, region, var_interanual * 100 AS v FROM v_variacion WHERE grupo = 'ipc_region'")
    # Desde la base del análisis hasta el último mes con RIPTE (la serie principal del gráfico).
    poder = db.consultar(
        con,
        "SELECT fecha, cbt_por_ripte AS ripte, cbt_por_smvm AS smvm, cbt_por_haber_minimo AS haber_minimo "
        "FROM v_poder_compra WHERE fecha >= :desde "
        "AND fecha <= (SELECT MAX(fecha) FROM v_poder_compra WHERE cbt_por_ripte IS NOT NULL) ORDER BY fecha",
        desde=config.INICIO_ANALISIS,
    )
    usd = db.consultar(
        con,
        "SELECT fecha, ripte_usd AS ripte, smvm_usd AS smvm, haber_minimo_usd AS haber_minimo "
        "FROM v_salario_usd WHERE fecha >= :desde AND ripte IS NOT NULL ORDER BY fecha",
        desde=config.INICIO_ANALISIS,
    )
    cobertura = q("SELECT serie_id, nombre, fuente, desde, hasta, meses, meses_de_rezago FROM v_cobertura")

    def serie_simple(df: pd.DataFrame, columnas: list[str], redondeo: int = 3) -> dict:
        return {
            "fechas": [f[:7] for f in df["fecha"]],
            "series": {c: [None if pd.isna(v) else round(float(v), redondeo) for v in df[c]] for c in columnas},
        }

    # El pico interanual sale de toda la serie, no de los 10 picos mensuales.
    pico_interanual = _registros(
        q(
            "SELECT etiqueta AS mes, ROUND(100 * var_interanual, 2) AS inflacion_interanual_pct FROM v_inflacion "
            "WHERE serie_id = 'ipc_nacional' AND ranking_interanual = 1 ORDER BY fecha LIMIT 1"
        )
    )[0]
    principal = consultas["01_respuesta_principal"]["filas"]
    por_sector = {f["serie_id"]: f for f in principal}
    ultima_inflacion = inflacion.iloc[-1]
    picos = consultas["03_inflacion_picos"]["filas"]
    poder_ultimo = next(f for f in consultas["06_poder_compra_canastas"]["filas"] if f["es_ultimo"])
    poder_inicio = consultas["06_poder_compra_canastas"]["filas"][0]
    usd_anual = consultas["08_salario_en_dolares"]["filas"]

    kpis = {
        "salario_real_total": por_sector["sal_total"],
        "salario_real_privado": por_sector["sal_privado"],
        "salario_real_publico": por_sector["sal_publico"],
        "salario_real_no_registrado": por_sector["sal_no_registrado"],
        "inflacion_mensual": {
            "mes": mes(inflacion.iloc[-1]["fecha"]),
            "valor": round(float(ultima_inflacion["mensual"]), 2),
        },
        "inflacion_interanual": {
            "mes": mes(inflacion.iloc[-1]["fecha"]),
            "valor": round(float(ultima_inflacion["interanual"]), 2),
        },
        "pico_mensual": picos[0],
        "pico_interanual": pico_interanual,
        "canastas_ripte": {"inicio": poder_inicio, "ultimo": poder_ultimo},
        "ripte_usd": {"inicio": usd_anual[0], "ultimo": usd_anual[-1]},
    }

    return {
        "meta": {
            "fecha_descarga": parametros["fecha_descarga"],
            "fecha_corte": parametros["fecha_corte"],
            "base_analisis": parametros["base_analisis"],
            "series": len(config.CATALOGO),
            "observaciones": con.execute("SELECT COUNT(*) FROM fact_valor").fetchone()[0],
        },
        "kpis": kpis,
        "sectores": NOMBRES_SECTOR,
        "regiones": config.REGIONES,
        "graficos": {
            "salario_real": _serie_ancha(real, "serie_id", "indice_real", 2),
            "inflacion": serie_simple(inflacion, ["mensual", "interanual"], 2),
            "categorias": _serie_ancha(categorias, "serie_id", "v", 2),
            "regiones": _serie_ancha(regiones, "region", "v", 2),
            "poder_compra": serie_simple(poder, ["ripte", "smvm", "haber_minimo"], 3),
            "usd": serie_simple(usd, ["ripte", "smvm", "haber_minimo"], 0),
        },
        "cobertura": _registros(cobertura),
        "consultas": consultas,
        "qa": reporte.a_dict(),
    }


def escribir_json(datos: dict, ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(datos, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def indices_calculadora(con: sqlite3.Connection) -> dict:
    return Indices.desde_base(con).a_dict()


# ---------------------------------------------------------------------------
# Power BI
# ---------------------------------------------------------------------------


def exportar_power_bi(con: sqlite3.Connection, destino: Path) -> list[Path]:
    """El modelo estrella en CSV (UTF-8 con BOM, separador coma, punto decimal)."""
    destino.mkdir(parents=True, exist_ok=True)
    tablas = {
        "dim_fecha": "SELECT fecha, anio, mes, trimestre, etiqueta FROM dim_fecha ORDER BY fecha",
        "dim_serie": "SELECT serie_id, nombre, grupo, region, unidad, base, fuente, desfase_meses FROM dim_serie",
        "fact_valor": "SELECT fecha, serie_id, valor, estado FROM fact_valor ORDER BY serie_id, fecha",
        "salario_real": "SELECT fecha, serie_id, indice_nominal, indice_real, var_real_interanual "
        "FROM v_salario_real ORDER BY serie_id, fecha",
    }
    rutas = []
    for nombre, sql in tablas.items():
        ruta = destino / f"{nombre}.csv"
        db.consultar(con, sql).to_csv(ruta, index=False, encoding="utf-8-sig")
        rutas.append(ruta)
    return rutas


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

AZUL = "1F4E79"
TITULOS_HOJA = {
    "01_respuesta_principal": "1 Respuesta",
    "02_ganadores_por_anio": "2 Sectores por año",
    "03_inflacion_picos": "3 Picos de inflación",
    "04_inflacion_anual_por_categoria": "4 Inflación por categoría",
    "05_rubros_ranking": "5 Rubros",
    "06_poder_compra_canastas": "6 Canastas",
    "07_regiones": "7 Regiones",
    "08_salario_en_dolares": "8 Salario en USD",
    "09_vigencia_de_conclusiones": "9 Vigencia",
}


def _alto_texto(texto: str, ancho_columnas: float) -> float:
    """Alto aproximado (en puntos) para que un texto con ajuste de línea se vea completo."""
    lineas = max(1, math.ceil(len(texto) * 1.1 / max(ancho_columnas, 1)))
    return 15 * lineas + 4


def _enlace(celda, hoja: str) -> None:
    """Hipervínculo interno a otra hoja del libro."""
    celda.hyperlink = Hyperlink(ref=celda.coordinate, location=f"'{hoja}'!A1", display=str(celda.value))
    celda.font = Font(color="0563C1", underline="single")


def _preparar_impresion(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _hoja_tabla(wb: Workbook, titulo: str, descripcion: str, df: pd.DataFrame, nombre_tabla: str):
    """Una hoja con título, la pregunta, cómo se lee y una tabla de Excel con formato."""
    df = df[[c for c in df.columns if c not in formato.OCULTAS]]
    ws = wb.create_sheet(titulo)
    for j, columna in enumerate(df.columns, start=1):
        valores = [formato.etiqueta(columna)] + [str(v) for v in df.iloc[:, j - 1].tolist()]
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(v) for v in valores) + 3, 45)
    ancho_total = sum(ws.column_dimensions[get_column_letter(j)].width for j in range(1, len(df.columns) + 1))
    columnas_texto = max(len(df.columns), 6)

    pregunta, _, lectura = descripcion.partition("Lectura:")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    for fila, texto, fuente in (
        (2, pregunta.removeprefix("Pregunta: ").strip(), Font(bold=True, size=11)),
        (3, f"Cómo se lee: {lectura.strip()}" if lectura else "", Font(italic=True, size=10, color="595959")),
    ):
        ws.cell(row=fila, column=1, value=texto or None).font = fuente
        ws.cell(row=fila, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=columnas_texto)
        ws.row_dimensions[fila].height = _alto_texto(texto, max(ancho_total, 60))

    fila0 = 5
    for j, columna in enumerate(df.columns, start=1):
        ws.cell(row=fila0, column=j, value=formato.etiqueta(columna))
    for i, fila in enumerate(df.itertuples(index=False), start=fila0 + 1):
        for j, valor in enumerate(fila, start=1):
            if isinstance(valor, str) and re.fullmatch(r"\d{4}-\d{2}-01", valor):
                valor = mes(valor)  # "2026-02-01" -> "feb-2026", como en el sitio
            ws.cell(row=i, column=j, value=_limpio(valor.item() if hasattr(valor, "item") else valor))
    ultima = fila0 + len(df)
    if len(df):
        ref = f"A{fila0}:{get_column_letter(len(df.columns))}{ultima}"
        tabla = Table(displayName=nombre_tabla, ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabla)
    for j in range(1, len(df.columns) + 1):
        if pd.api.types.is_float_dtype(df.iloc[:, j - 1]):
            for i in range(fila0 + 1, ultima + 1):
                ws.cell(row=i, column=j).number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
    _enlace(ws.cell(row=ultima + 2, column=1, value="← Volver al índice"), "Leeme")
    _preparar_impresion(ws)
    return ws


def exportar_excel(con: sqlite3.Connection, consultas: dict, reporte: Reporte, ruta: Path) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.properties.creator = "Juan Barros"
    wb.properties.title = "¿Le ganaron los salarios a la inflación? Argentina 2016–2026"
    portada = wb.active
    portada.title = "Leeme"
    parametros = dict(con.execute("SELECT clave, valor FROM parametro").fetchall())
    descarga = parametros["fecha_descarga"]
    portada["A1"] = "¿Le ganaron los salarios a la inflación? Argentina 2016–2026"
    portada["A1"].font = Font(bold=True, size=16, color=AZUL)
    portada["A2"] = (
        "Datos oficiales de INDEC, BCRA, ANSES y MTEySS descargados el "
        f"{descarga[8:10]}/{descarga[5:7]}/{descarga[:4]}. "
        "Cada hoja numerada responde una pregunta de negocio con el resultado de una consulta SQL de sql/consultas."
    )
    portada["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    portada.merge_cells("A2:B2")
    portada.row_dimensions[2].height = 34
    portada["A4"], portada["B4"] = "Hoja", "Qué responde"
    for celda in (portada["A4"], portada["B4"]):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
    indice = [
        (
            TITULOS_HOJA.get(nombre, nombre[:31]),
            datos["descripcion"].partition("Lectura:")[0].removeprefix("Pregunta: "),
        )
        for nombre, datos in consultas.items()
    ]
    indice += [
        ("Salario real mensual", "La serie mensual de salario real por sector, con gráfico."),
        ("QA", "Los chequeos automáticos de calidad de datos y su resultado."),
    ]
    for i, (hoja, que) in enumerate(indice, start=5):
        _enlace(portada.cell(row=i, column=1, value=hoja), hoja)
        portada.cell(row=i, column=1).alignment = Alignment(vertical="top")
        portada.cell(row=i, column=2, value=que.strip()).alignment = Alignment(wrap_text=True, vertical="top")
    fila_autor = 6 + len(indice)
    portada.cell(row=fila_autor, column=1, value="Autor: Juan Barros · github.com/JuanMBF/InflationArg")
    portada.cell(row=fila_autor, column=1).font = Font(italic=True, color="595959")
    portada.column_dimensions["A"].width = 26
    portada.column_dimensions["B"].width = 95
    _preparar_impresion(portada)

    for nombre, datos in consultas.items():
        df = pd.DataFrame(datos["filas"], columns=datos["columnas"])
        _hoja_tabla(wb, TITULOS_HOJA.get(nombre, nombre[:31]), datos["descripcion"], df, f"t_{nombre[:3]}")

    real = db.consultar(con, "SELECT fecha, serie_id, indice_real FROM v_salario_real")
    ancho = real.pivot(index="fecha", columns="serie_id", values="indice_real").sort_index()
    ancho = ancho[[c for c in NOMBRES_SECTOR if c in ancho.columns]]
    ancho.columns = [NOMBRES_SECTOR[c] for c in ancho.columns]
    ancho.index = [mes(f) for f in ancho.index]
    tabla = ancho.round(2).reset_index(names="mes")
    ws = _hoja_tabla(
        wb,
        "Salario real mensual",
        "Pregunta: ¿Cómo evolucionó mes a mes el salario real de cada sector? "
        "Lectura: índice de salarios del INDEC deflactado por el IPC; 100 = poder de compra de dic-2016.",
        tabla,
        "t_salario_real",
    )
    grafico = LineChart()
    grafico.title = "Salario real por sector (dic-2016 = 100)"
    grafico.y_axis.title = "Índice"
    grafico.height, grafico.width = 9, 22
    datos = Reference(ws, min_col=2, max_col=1 + len(ancho.columns), min_row=5, max_row=5 + len(tabla))
    grafico.add_data(datos, titles_from_data=True)
    grafico.set_categories(Reference(ws, min_col=1, min_row=6, max_row=5 + len(tabla)))
    # openpyxl 3.1 marca los ejes como borrados: sin esto Excel no muestra fechas ni valores.
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.y_axis.numFmt = "0"
    grafico.y_axis.scaling.min = math.floor(float(ancho.min().min()) / 10) * 10
    grafico.x_axis.tickLblSkip = 12
    grafico.x_axis.tickMarkSkip = 12
    for serie, (color, guion) in zip(
        grafico.series, [("2A78D6", None), ("EB6834", None), ("1BAF7A", None), ("121211", "dash")], strict=False
    ):
        serie.smooth = False
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 22000
        if guion:
            serie.graphicalProperties.line.dashStyle = guion
    ws.add_chart(grafico, "H5")

    qa_df = pd.DataFrame([{"nivel": r.nivel, "chequeo": r.titulo, "resultado": r.detalle} for r in reporte.resultados])
    ws = _hoja_tabla(
        wb,
        "QA",
        "Pregunta: ¿Los datos son confiables? Lectura: un error frena el pipeline; una alerta es una anomalía "
        "real de las fuentes que se documenta y, si corresponde, se aparta del análisis.",
        qa_df,
        "t_qa",
    )
    colores = {"ok": "C6EFCE", "info": "DDEBF7", "alerta": "FFEB9C", "error": "FFC7CE"}
    ws.column_dimensions["C"].width = 100
    for i, r in enumerate(reporte.resultados, start=6):
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=colores[r.nivel])
        ws.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = _alto_texto(r.detalle, 100)

    wb.save(ruta)
    return ruta
